from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from copy import deepcopy
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from config import DOCUMENTS_DIR, TEMPLATE_PATH
from utils import clean_rank, format_date_id, slugify


class DocumentGenerationError(RuntimeError):
    pass


BODY_FONT = "Arial"
BODY_FONT_SIZE = Pt(12)
SIGNATURE_TEXT_LEFT_INDENT = Pt(6.75)


def _remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _copy_paragraph_properties(source: Paragraph, target: Paragraph) -> None:
    source_ppr = source._p.pPr
    target_ppr = target._p.pPr
    if target_ppr is not None:
        target._p.remove(target_ppr)
    if source_ppr is not None:
        target._p.insert(0, deepcopy(source_ppr))


def _set_run_font(run, *, bold: bool | None = None, underline: bool | None = None) -> None:
    run.font.name = BODY_FONT
    run.font.size = BODY_FONT_SIZE
    if bold is not None:
        run.bold = bold
    if underline is not None:
        run.underline = underline
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), BODY_FONT)


def _format_paragraph(
    paragraph: Paragraph,
    *,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
    bold: bool | None = None,
    keep_with_next: bool | None = None,
    space_before: float | None = None,
    space_after: float | None = None,
    line_spacing: float | None = None,
) -> None:
    """Format only the properties explicitly requested.

    IMPORTANT: blank paragraphs and paragraph spacing from the official template are
    layout elements. Do not normalize them to zero. This keeps intentional empty
    space exactly as it exists in the template.
    """
    if alignment is not None:
        paragraph.alignment = alignment
    if space_before is not None:
        paragraph.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        paragraph.paragraph_format.space_after = Pt(space_after)
    if line_spacing is not None:
        paragraph.paragraph_format.line_spacing = line_spacing
    if keep_with_next is not None:
        paragraph.paragraph_format.keep_with_next = keep_with_next
    for run in paragraph.runs:
        _set_run_font(run, bold=bold)


def _set_paragraph_text(paragraph: Paragraph, text: str, *, bold: bool | None = None) -> None:
    # paragraph.text keeps paragraph-level tabs/indents from the template.
    paragraph.text = text
    for run in paragraph.runs:
        _set_run_font(run, bold=bold)


def _set_paragraph_text_preserve_objects(
    paragraph: Paragraph, text: str, *, bold: bool | None = None
) -> None:
    """Replace only direct paragraph text while preserving drawings/text boxes.

    Some official-template paragraphs also anchor floating signature text boxes.
    Assigning ``paragraph.text`` would delete those objects. This helper edits only
    direct ``w:t`` nodes in normal runs and leaves drawing/XML objects untouched.
    """
    direct_text_nodes = []
    for run_el in paragraph._p.findall(qn("w:r")):
        direct_text_nodes.extend(run_el.findall(qn("w:t")))

    if direct_text_nodes:
        direct_text_nodes[0].text = text
        if text.startswith(" ") or text.endswith(" ") or "  " in text:
            direct_text_nodes[0].set(qn("xml:space"), "preserve")
        for node in direct_text_nodes[1:]:
            node.text = ""
    else:
        run = paragraph.add_run(text)
        _set_run_font(run, bold=bold)

    for run in paragraph.runs:
        _set_run_font(run, bold=bold)


def _remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _set_cell_margins(cell, *, top: int = 0, start: int = 0, bottom: int = 0, end: int = 0) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _clear_cell(cell) -> None:
    cell.text = ""
    # python-docx leaves one empty paragraph, which is what we want.


def _add_cell_line(
    cell,
    text: str,
    *,
    bold: bool = False,
    underline: bool = False,
    keep_with_next: bool = False,
    space_after: float = 0,
) -> Paragraph:
    # Reuse the single empty paragraph first. ``cell.text = ""`` may leave an
    # empty run, so text emptiness is the reliable check. This keeps generated
    # employee rows compact without touching blank space elsewhere in the template.
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, bold=bold, underline=underline)
    _format_paragraph(
        p,
        keep_with_next=keep_with_next,
        space_before=0,
        space_after=space_after,
        line_spacing=1.0,
    )
    return p


def _replace_legal_bases(doc: Document, legal_bases: list[str]) -> None:
    paragraphs = doc.paragraphs
    mem_index = next(
        i for i, p in enumerate(paragraphs) if p.text.strip() == "MEMERINTAHKAN"
    )
    base_indices = [
        i
        for i, p in enumerate(paragraphs[:mem_index])
        if p.text.strip().startswith("DASAR")
        or (p.text.strip() and p.text.lstrip().split(".", 1)[0].isdigit())
    ]
    base_indices = [i for i in base_indices if 0 < i < mem_index]
    if not base_indices:
        raise DocumentGenerationError("Blok DASAR tidak ditemukan pada template.")

    first_template = paragraphs[base_indices[0]]
    continuation_template = paragraphs[base_indices[1]] if len(base_indices) > 1 else first_template
    mem_paragraph = paragraphs[mem_index]

    for index in reversed(base_indices):
        _remove_paragraph(doc.paragraphs[index])

    for number, text in enumerate(legal_bases, start=1):
        new_p = mem_paragraph.insert_paragraph_before()
        _copy_paragraph_properties(first_template if number == 1 else continuation_template, new_p)
        prefix = f"DASAR\t:\t{number}.\t" if number == 1 else f"\t\t{number}.\t"
        _set_paragraph_text(new_p, prefix + text.strip())
        _format_paragraph(new_p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

    _format_paragraph(mem_paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)


def _ensure_blank_line_after_memerintahkan(doc: Document) -> None:
    """Ensure one intentional empty paragraph immediately after MEMERINTAHKAN.

    The official layout requires a visible blank line before the KEPADA block.
    Existing blank paragraphs are never removed or normalized.
    """
    mem_paragraph = next(
        (p for p in doc.paragraphs if p.text.strip() == "MEMERINTAHKAN"),
        None,
    )
    if mem_paragraph is None:
        raise DocumentGenerationError("Bagian MEMERINTAHKAN tidak ditemukan pada template.")

    next_element = mem_paragraph._p.getnext()
    if next_element is not None and next_element.tag == qn("w:p"):
        text = "".join(node.text or "" for node in next_element.findall(".//" + qn("w:t")))
        if not text.strip():
            return

    # Reuse a real empty paragraph from the official template so its spacing
    # characteristics are preserved instead of inventing a new layout style.
    blank_template = next((p for p in reversed(doc.paragraphs) if not p.text.strip()), None)
    blank_element = deepcopy(blank_template._p) if blank_template is not None else OxmlElement("w:p")
    mem_paragraph._p.addnext(blank_element)


def _set_table_indent(table, inches: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(int(inches * 1440)))
    tbl_ind.set(qn("w:type"), "dxa")


def _replace_employee_table(doc: Document, employees: list[dict[str, Any]]) -> None:
    if not doc.tables:
        raise DocumentGenerationError("Tabel KEPADA tidak ditemukan pada template.")

    # Preserve the official KEPADA block exactly as authored in the template.
    # Only the old floating employee content in the right cell is cleared.
    kepada_table = doc.tables[0]
    if not kepada_table.rows or len(kepada_table.columns) < 2:
        raise DocumentGenerationError("Format tabel KEPADA tidak sesuai.")

    kepada_left = kepada_table.rows[0].cells[0]
    kepada_right = kepada_table.rows[0].cells[1]
    _clear_cell(kepada_right)
    _set_row_cant_split(kepada_table.rows[0])

    # Do not rewrite the word KEPADA or its tabs/colon. Keep the original
    # paragraph properties verbatim and only normalize the requested font.
    for paragraph in kepada_left.paragraphs:
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            _set_run_font(run)

    # Dynamic employee list starts below KEPADA. This keeps the official KEPADA
    # line unchanged while still allowing many employees to flow across pages.
    employee_table = doc.add_table(rows=0, cols=4)
    employee_table.autofit = False
    _remove_table_borders(employee_table)
    _set_table_indent(employee_table, 1.00)

    tbl_layout = employee_table._tbl.tblPr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        employee_table._tbl.tblPr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    # Widths are the employee portion of the previous five-column layout.
    # Total width = 5.68 in, positioned 1.00 in from the left text margin.
    widths = [Inches(0.38), Inches(1.20), Inches(0.22), Inches(3.88)]
    for column, width in zip(employee_table.columns, widths):
        column.width = width

    for index, employee in enumerate(employees, start=1):
        row = employee_table.add_row()
        _set_row_cant_split(row)
        values = [
            ("Nama", str(employee.get("name", "")).strip()),
            ("Pangkat/gol", clean_rank(str(employee.get("rank", "")))),
            ("NIP", str(employee.get("identifier", "")).strip()),
            ("Jabatan", str(employee.get("position", "")).strip()),
        ]

        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_cell_margins(cell, top=0, start=0, bottom=0, end=35)
            _clear_cell(cell)

        _add_cell_line(row.cells[0], f"{index}.")

        for line_idx, (label, value) in enumerate(values):
            keep = line_idx < len(values) - 1
            spacing = 8 if line_idx == len(values) - 1 else 0
            _add_cell_line(row.cells[1], label, keep_with_next=keep, space_after=spacing)
            _add_cell_line(row.cells[2], ":", keep_with_next=keep, space_after=spacing)
            _add_cell_line(row.cells[3], value, keep_with_next=keep, space_after=spacing)

        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _format_paragraph(paragraph)

    # Keep the official KEPADA table in place. Insert one template-style blank
    # paragraph after it, then the dynamic employee table. This mirrors the
    # spacing of the official two-person template and does not alter KEPADA.
    blank_template = next((p for p in reversed(doc.paragraphs) if not p.text.strip()), None)
    spacer = deepcopy(blank_template._p) if blank_template is not None else OxmlElement("w:p")
    kepada_table._tbl.addnext(spacer)
    spacer.addnext(employee_table._tbl)


def _replace_number(doc: Document, full_number: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("Nomor"):
            _set_paragraph_text(paragraph, f"Nomor : {full_number}")
            _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            return
    raise DocumentGenerationError("Baris Nomor surat tidak ditemukan pada template.")


def _replace_purpose(doc: Document, purpose_text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("UNTUK"):
            _set_paragraph_text(paragraph, f"UNTUK\t:\t{purpose_text.strip()}")
            # Preserve the template's original paragraph spacing/blank space.
            _format_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
            return
    raise DocumentGenerationError("Bagian UNTUK tidak ditemukan pada template.")


def _replace_signature_block(doc: Document, issue_date: date, issue_city: str) -> None:
    """Update only city/date text and preserve the signature layout verbatim.

    The official template can contain intentional blank paragraphs, spacing, and
    floating objects/text boxes for the signatory. Those must not be cleared or
    rebuilt because doing so changes the template's empty space and page balance.
    """
    if len(doc.tables) < 2:
        raise DocumentGenerationError("Blok penetapan surat tidak ditemukan pada template.")

    table = doc.tables[-1]
    if not table.rows or len(table.columns) < 2:
        raise DocumentGenerationError("Format blok penetapan surat tidak sesuai.")

    right = table.rows[0].cells[1]
    found_city = False
    found_date = False
    for paragraph in right.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Ditetapkan di"):
            _set_paragraph_text_preserve_objects(paragraph, f"Ditetapkan di {issue_city}")
            paragraph.paragraph_format.left_indent = SIGNATURE_TEXT_LEFT_INDENT
            found_city = True
        elif text.startswith("pada tanggal"):
            _set_paragraph_text_preserve_objects(paragraph, f"pada tanggal {format_date_id(issue_date)}")
            paragraph.paragraph_format.left_indent = SIGNATURE_TEXT_LEFT_INDENT
            found_date = True

    # Fallback only when a custom template is missing one of the two placeholders.
    # Existing empty paragraphs are reused first so no intentional blank space is deleted.
    if not found_city:
        target = next((p for p in right.paragraphs if not p.text.strip()), None)
        target = target or right.add_paragraph()
        _set_paragraph_text(target, f"Ditetapkan di {issue_city}")
        target.paragraph_format.left_indent = SIGNATURE_TEXT_LEFT_INDENT
    if not found_date:
        target = next((p for p in right.paragraphs if not p.text.strip()), None)
        target = target or right.add_paragraph()
        _set_paragraph_text(target, f"pada tanggal {format_date_id(issue_date)}")
        target.paragraph_format.left_indent = SIGNATURE_TEXT_LEFT_INDENT

    # Font is normalized without touching paragraph spacing, empty paragraphs,
    # line spacing, tabs, indents, or floating signature objects.
    for paragraph in right.paragraphs:
        for run in paragraph.runs:
            _set_run_font(run)


def _normalize_body_font(doc: Document) -> None:
    """Set body text to Arial 12 without changing template layout.

    Do not alter paragraph spacing, line spacing, blank paragraphs, tabs, indents,
    keep settings, or other geometry. Empty space in the template is intentional.
    """
    body_started = False
    for p in doc.paragraphs:
        if p.text.strip() == "SURAT TUGAS":
            body_started = True
        if body_started:
            for run in p.runs:
                _set_run_font(run, bold=True if p.text.strip() == "SURAT TUGAS" else None)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        _set_run_font(run)
                for nested in cell.tables:
                    for nrow in nested.rows:
                        for ncell in nrow.cells:
                            for p in ncell.paragraphs:
                                for run in p.runs:
                                    _set_run_font(run)


def generate_docx(
    *,
    full_number: str,
    sequence_number: int,
    destination: str,
    activity: str,
    issue_date: date,
    employees: list[dict[str, Any]],
    legal_bases: list[str],
    purpose_text: str,
    version: int = 1,
    output_dir: Path = DOCUMENTS_DIR,
    issue_city: str = "Surabaya",
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    if not employees:
        raise DocumentGenerationError("Minimal satu pegawai harus dipilih.")
    if not template_path.exists():
        raise DocumentGenerationError(f"Template tidak ditemukan: {template_path}")
    if len(legal_bases) < 3:
        raise DocumentGenerationError("Minimal tiga dasar utama harus tersedia.")

    output_dir.mkdir(parents=True, exist_ok=True)
    version_suffix = "" if version <= 1 else f"_v{version}"
    number_part = str(sequence_number) if sequence_number > 0 else "TANPA_NOMOR"
    filename = (
        f"SPT_{number_part}_{slugify(destination, 35)}_"
        f"{slugify(activity, 45)}{version_suffix}.docx"
    )
    output_path = output_dir / filename
    shutil.copy2(template_path, output_path)

    doc = Document(output_path)
    _replace_number(doc, full_number)
    _replace_legal_bases(doc, legal_bases)
    _ensure_blank_line_after_memerintahkan(doc)
    _replace_employee_table(doc, employees)
    _replace_purpose(doc, purpose_text)
    _replace_signature_block(doc, issue_date, issue_city)
    _normalize_body_font(doc)
    doc.save(output_path)
    return output_path


def convert_docx_to_pdf(docx_path: Path) -> Path | None:
    """Convert DOCX to PDF using LibreOffice. Returns None when unavailable."""
    docx_path = Path(docx_path)
    if not docx_path.exists():
        raise FileNotFoundError(docx_path)

    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return None

    output_dir = docx_path.parent
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile_dir:
        env = os.environ.copy()
        env["HOME"] = profile_dir
        result = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(docx_path),
            ],
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=90,
        )
    pdf_path = docx_path.with_suffix(".pdf")
    if result.returncode != 0 or not pdf_path.exists() or pdf_path.stat().st_size == 0:
        return None
    return pdf_path
